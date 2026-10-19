"""Builds the .docx meeting report: a one-page brief first, the detail and the transcript behind it."""

import io
import re
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

from .config import reports_dir  # noqa: E402

PALETTE = ["#2563eb", "#16a34a", "#f59e0b", "#dc2626", "#7c3aed", "#0891b2", "#db2777", "#65a30d"]
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
MERGE_GAP = 8.0  # seconds: consecutive lines by the same speaker closer than this read as one paragraph


# ------------------------------------------------------------------ text helpers (also unit-tested)
def stamp(t: float) -> str:
    t = int(t or 0)
    return f"{t // 60:02d}:{t % 60:02d}"


def _norm(s: str) -> str:
    return re.sub(r"\W+", " ", s or "", flags=re.UNICODE).lower().strip()


def same_text(a: str, b: str) -> bool:
    return _norm(a) == _norm(b)


def similar_text(a: str, b: str, threshold: float = 0.8) -> bool:
    """True when two texts are the same words give or take speech-to-text wobble (an English line and its 'translation')."""
    wa, wb = set(_norm(a).split()), set(_norm(b).split())
    if not wa or not wb:
        return wa == wb
    return len(wa & wb) / max(len(wa), len(wb)) >= threshold


def first_sentence(text: str, limit: int = 220) -> str:
    text = (text or "").strip()
    m = re.match(r"(.+?[.!?])(\s|$)", text, flags=re.S)
    out = m.group(1) if m else text
    return out if len(out) <= limit else out[: limit - 1].rstrip() + "..."


def merge_transcript(lines, gap: float = MERGE_GAP) -> list[dict]:
    """One paragraph per speaker turn instead of one per phrase; drops translations that merely repeat the text."""
    merged: list[dict] = []
    for x in sorted(lines, key=lambda l: l.get("t", 0)):
        text = (x.get("text") or "").strip()
        tr = (x.get("translation") or "").strip()
        if not text and not tr:
            continue
        prev = merged[-1] if merged else None
        if prev and prev["speaker"] == x.get("speaker") and x.get("t", 0) - prev["last_t"] <= gap:
            prev["text"] = f"{prev['text']} {text}".strip()
            prev["translation"] = f"{prev['translation']} {tr}".strip()
            prev["last_t"] = x.get("t", 0)
            prev["mention"] = prev["mention"] or bool(x.get("mention"))
        else:
            merged.append({"t": x.get("t", 0), "last_t": x.get("t", 0), "speaker": x.get("speaker") or "Speaker",
                           "text": text, "translation": tr, "mention": bool(x.get("mention"))})
    for m in merged:
        if not m["text"]:  # only a translation exists
            m["text"], m["translation"] = m["translation"], ""
        elif similar_text(m["text"], m["translation"]):
            m["translation"] = ""
    return merged


def bookmark_context(bookmarks, transcript, within: float = 30.0) -> list[dict]:
    """Give each bookmark the words nearest to it, so the report says WHY the moment mattered."""
    out = []
    for b in bookmarks or []:
        b = dict(b)
        if not (b.get("text") or "").strip() and transcript:
            near = min(transcript, key=lambda l: abs(l.get("t", 0) - b.get("t", 0)))
            if abs(near.get("t", 0) - b.get("t", 0)) <= within:
                b["text"] = ((near.get("translation") or near.get("text") or "").strip())[:160]
        out.append(b)
    return out


def speaker_totals(talk) -> dict[str, float]:
    return {s: sum(per.values()) for s, per in (talk or {}).items() if sum(per.values()) > 0}


def real_speakers(totals: dict[str, float]) -> dict[str, float]:
    """Speakers worth charting: a single bar or 'Others 100%' tells the reader nothing."""
    named = {s: v for s, v in totals.items() if s.lower() != "others"}
    return totals if len(totals) >= 2 and len(named) >= 1 else {}


# ------------------------------------------------------------------ docx helpers
TCPR_ORDER = ["tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
              "tcFitText", "vAlign", "hideMark"]


def _tcpr_put(tcPr, el) -> None:
    """Add/replace a child of <w:tcPr> and keep children in schema order (Word rejects other orders)."""
    for old in tcPr.findall(el.tag):
        tcPr.remove(old)
    tcPr.append(el)
    kids = sorted(tcPr, key=lambda k: TCPR_ORDER.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in TCPR_ORDER else 99)
    for k in kids:
        tcPr.remove(k)
    for k in kids:
        tcPr.append(k)


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _tcpr_put(tcPr, shd)


def _borders(cell, **edges) -> None:
    """edges: left=("2563EB", 24), ... (colour, eighths of a point) or None for no line."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:color"), spec[0])
            el.set(qn("w:sz"), str(spec[1]))
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    _tcpr_put(tcPr, b)


def _cell_margins(cell, top=90, bottom=90, left=160, right=160) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    _tcpr_put(tcPr, m)


def _run(par, text, size=11, bold=False, italic=False, color=INK):
    r = par.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return r


def _space(par, before=0, after=4, line=None):
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        pf.line_spacing = line


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    _space(p, before=16, after=5)
    p.paragraph_format.keep_with_next = True
    _run(p, text, size=13.5, bold=True)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "2"), ("color", "E2E8F0")):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.insert_element_before(bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                              "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                              "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                              "w:pPrChange")


def _bullet(doc, text: str, bold_lead: str = ""):
    p = doc.add_paragraph(style="List Bullet")
    _space(p, after=3)
    if bold_lead:
        _run(p, bold_lead, bold=True)
    _run(p, text)
    return p


def _callout(doc, label: str, body: str, fill="EFF6FF", bar="2563EB", size=13):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    _shade(c, fill)
    _borders(c, left=(bar, 36))
    _cell_margins(c, 120, 120, 200, 200)
    p = c.paragraphs[0]
    _space(p, after=2)
    _run(p, label.upper(), size=8.5, bold=True, color=ACCENT if bar == "2563EB" else MUTED)
    p2 = c.add_paragraph()
    _space(p2, after=0, line=1.15)
    _run(p2, body, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _table(doc, header: list[str], rows: list[list[str]], widths: list[float]):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _shade(c, "F1F5F9")
        _borders(c, bottom=("CBD5E1", 8))
        _cell_margins(c, 70, 70, 120, 120)
        _run(c.paragraphs[0], h.upper(), size=8.5, bold=True, color=MUTED)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            _borders(cells[i], bottom=("E2E8F0", 4))
            _cell_margins(cells[i], 80, 80, 120, 120)
            _run(cells[i].paragraphs[0], v, size=10.5)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    return t


def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight", transparent=False, facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def _share_chart(totals: dict[str, float]):
    fig, ax = plt.subplots(figsize=(3.6, 2.6))
    names = list(totals)
    ax.pie([totals[n] for n in names], labels=names, colors=PALETTE[: len(names)], autopct="%1.0f%%", startangle=90,
           wedgeprops={"width": 0.45, "edgecolor": "white"}, textprops={"fontsize": 9})
    return _png(fig)


# ------------------------------------------------------------------ the report
def build(*, settings, title, platform, started: datetime, minutes: float, transcript, talk, analysis,
          languages, meetings_total: int, minutes_total: float, participants=(), bookmarks=()) -> Path:
    analysis = analysis or {}
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    host = (getattr(settings, "name", "") or "").strip()
    totals = speaker_totals(talk)
    heading = title or analysis.get("title") or "Meeting report"

    # ---- header: title + one quiet line of facts (only the facts we actually have)
    p = doc.add_paragraph()
    _space(p, after=2)
    _run(p, heading, size=24, bold=True)
    others = [n for n in dict.fromkeys(list(participants) + list(totals)) if n.lower() not in ("others", "you")
              and (not host or host.lower() not in n.lower())]
    facts = [f"{started:%a %d %b %Y, %H:%M}", f"{max(minutes, 0.1):.0f} min" if minutes >= 1 else f"{minutes * 60:.0f} sec"]
    if platform:
        facts.append(platform)
    if others:
        facts.append("with " + ", ".join(others[:6]))
    if analysis.get("sentiment"):
        facts.append(f"tone: {str(analysis['sentiment']).lower()}")
    p = doc.add_paragraph()
    _space(p, after=10)
    _run(p, "  ·  ".join(facts), size=10, color=MUTED)

    # ---- the whole point, in one sentence
    tldr = (analysis.get("tldr") or first_sentence(analysis.get("summary", ""))).strip()
    if tldr:
        _callout(doc, "In short", tldr)

    # ---- what happened and what to do
    decisions = analysis.get("decisions") or []
    if decisions:
        _heading(doc, "Decisions")
        for d in decisions:
            p = doc.add_paragraph()
            _space(p, after=3)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            _run(p, "✓  ", bold=True, color=RGBColor(0x16, 0xA3, 0x4A))
            _run(p, str(d))

    actions = analysis.get("action_items") or []
    if actions:
        _heading(doc, "Action items")
        has_due = any((a.get("due") or "").strip() for a in actions)
        header = ["Task", "Who"] + (["When"] if has_due else [])
        widths = [4.0, 1.4] + ([1.3] if has_due else [])
        if not has_due:
            widths = [4.9, 1.8]
        rows = [[str(a.get("task", "")), str(a.get("owner") or "-")] + ([str(a.get("due") or "-")] if has_due else [])
                for a in actions]
        _table(doc, header, rows, widths)

    if analysis.get("summary") and tldr and not same_text(tldr, analysis["summary"]):
        _heading(doc, "Context")
        p = doc.add_paragraph()
        _space(p, after=4, line=1.15)
        _run(p, analysis["summary"])

    takeaways = analysis.get("takeaways") or analysis.get("key_points") or []
    if takeaways:
        _heading(doc, "Worth remembering")
        for k in takeaways:
            _bullet(doc, str(k))

    if analysis.get("open_questions"):
        _heading(doc, "Still open")
        for q in analysis["open_questions"]:
            _bullet(doc, str(q))

    # ---- detail
    chapters = analysis.get("chapters") or []
    if chapters:
        _heading(doc, "Timeline")
        rows = [[str(c.get("start", "")), f"{c.get('title', '')}" + (f" – {c['summary']}" if c.get("summary") else "")]
                for c in chapters]
        _table(doc, ["Time", "What happened"], rows, [0.8, 5.9])

    bookmarks = bookmark_context(bookmarks, transcript)
    mentions = [x for x in transcript if x.get("mention")]
    if bookmarks or mentions:
        _heading(doc, "Highlights")
        for b in bookmarks:
            _bullet(doc, b.get("text") or "You bookmarked this moment.", f"★ {stamp(b.get('t', 0))}  ")
        for x in mentions:
            _bullet(doc, x.get("translation") or x.get("text") or "", f"@ {stamp(x.get('t', 0))}  {x['speaker']} mentioned you: ")

    speakers = real_speakers(totals)
    if speakers:
        _heading(doc, "Who spoke")
        doc.add_picture(_share_chart(speakers), width=Inches(2.6))
        whole = sum(speakers.values())
        _table(doc, ["Speaker", "Minutes", "Share"],
               [[s, f"{v / 60:.1f}", f"{v / whole:.0%}"] for s, v in sorted(speakers.items(), key=lambda kv: -kv[1])],
               [3.4, 1.6, 1.6])

    email = analysis.get("follow_up_email") or {}
    if email.get("body"):
        _heading(doc, "Follow-up email (draft)")
        t = doc.add_table(rows=1, cols=1)
        c = t.rows[0].cells[0]
        _shade(c, "F8FAFC")
        _borders(c, top=("E2E8F0", 6), left=("E2E8F0", 6), bottom=("E2E8F0", 6), right=("E2E8F0", 6))
        _cell_margins(c, 140, 140, 200, 200)
        p = c.paragraphs[0]
        _space(p, after=6)
        _run(p, "Subject: ", size=10.5, bold=True, color=MUTED)
        _run(p, str(email.get("subject", "")), size=10.5, bold=True)
        for para in str(email["body"]).split("\n"):
            q = c.add_paragraph()
            _space(q, after=0 if para.strip() else 2, line=1.1)
            _run(q, para, size=10.5)

    p = doc.add_paragraph()
    _space(p, before=14)
    _run(p, f"Ink has recorded {meetings_total} meeting{'s' if meetings_total != 1 else ''} for you "
            f"({minutes_total:.0f} min in total).", size=9, color=MUTED)

    # ---- the transcript, readable: one paragraph per turn, no repeated lines
    if transcript:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _heading(doc, "Full transcript")
        langs = ", ".join(languages)
        if langs:
            p = doc.add_paragraph()
            _run(p, f"Languages heard: {langs}. English is shown under any line that was spoken in another language.",
                 size=9.5, color=MUTED)
        for m in merge_transcript(transcript):
            p = doc.add_paragraph()
            _space(p, before=6, after=1, line=1.1)
            _run(p, f"{stamp(m['t'])}  ", size=9, color=MUTED)
            _run(p, f"{m['speaker']}  ", size=10.5, bold=True, color=ACCENT)
            if m["mention"]:
                _run(p, "@you  ", size=9, bold=True, color=RGBColor(0x7C, 0x3A, 0xED))
            _run(p, m["text"], size=10.5)
            if m["translation"]:
                q = doc.add_paragraph()
                _space(q, after=2, line=1.1)
                q.paragraph_format.left_indent = Inches(0.45)
                _run(q, "In English: ", size=9.5, bold=True, color=MUTED)
                _run(q, m["translation"], size=10, italic=True, color=MUTED)

    out = reports_dir(settings.name)
    out.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[<>:"/\\|?*\r\n]', "", title).strip(" .")[:80]
    stamp_name = f"{started:%Y-%m-%d %H-%M}"
    path = out / (f"{safe} - {stamp_name}.docx" if safe else f"{stamp_name}.docx")
    doc.save(path)
    return path
